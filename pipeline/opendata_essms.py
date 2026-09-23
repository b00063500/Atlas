"""Open data ESSMS : le xlsx de la HAS contient TOUT — identité des
établissements, dates et cotations des évaluations. Pas de fiche à crawler,
pas de PDF : ce module télécharge le tableau et le rend ligne par ligne.
Module pur : il renvoie des données, il ne touche jamais la base."""
import datetime

import openpyxl

from config import INDEX_ESSMS, URL_INDEX_ESSMS, rafraichir


def telecharger_index() -> bool:
    """Rafraîchit le xlsx local. Renvoie True si une nouvelle version a été
    écrite, False si le serveur a répondu 304 (« tu as déjà la bonne »).
    La mécanique du GET conditionnel vit dans config.rafraichir."""
    return rafraichir(URL_INDEX_ESSMS, INDEX_ESSMS, timeout=600)


def lire_evaluations(colonnes_attendues: tuple[str, ...]):
    """Générateur : un dict {colonne: valeur} par ligne du tableau.

    Vérifie AVANT de rendre la première ligne que toutes les colonnes
    attendues par le schéma existent encore : si la HAS renomme une colonne,
    on veut mourir ici avec son nom, pas insérer des NULL en silence."""
    classeur = openpyxl.load_workbook(INDEX_ESSMS, read_only=True, data_only=True)
    try:
        feuille = classeur[classeur.sheetnames[0]]
        lignes = feuille.iter_rows(values_only=True)
        entete = [str(nom) if nom is not None else "" for nom in next(lignes)]

        disparues = set(colonnes_attendues) - set(entete)
        if disparues:
            raise RuntimeError(f"colonnes absentes du xlsx HAS : {sorted(disparues)}")

        for ligne in lignes:
            yield {nom: _convertir(valeur) for nom, valeur in zip(entete, ligne)}
    finally:
        classeur.close()


def _convertir(valeur):
    """openpyxl → SQLite : datetime en date ISO, booléen en 0/1,
    chaîne vide en NULL. Les nombres passent tels quels."""
    if isinstance(valeur, datetime.datetime):
        return valeur.date().isoformat()
    if isinstance(valeur, bool):
        return int(valeur)
    if isinstance(valeur, str):
        valeur = valeur.strip()
        return valeur if valeur else None
    return valeur
