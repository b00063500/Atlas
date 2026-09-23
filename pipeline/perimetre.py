"""Open data : du xlsx Qualiscope à la liste des FINESS certifiés.
Ne concerne que le sanitaire — l'ESSMS n'a pas de périmètre à crawler,
ses évaluations arrivent entières par l'open data (opendata_essms.py)."""
import openpyxl

from config import INDEX_SANITAIRE, URL_INDEX_SANITAIRE, rafraichir


def telecharger_index() -> bool:
    """Rafraîchit le xlsx Qualiscope (GET conditionnel, voir config.rafraichir).
    Un établissement nouvellement certifié n'entre dans le périmètre que si
    ce fichier est remis à jour : il l'est à chaque tournée, pas seulement
    la première fois."""
    return rafraichir(URL_INDEX_SANITAIRE, INDEX_SANITAIRE, timeout=300)


def perimetre_sanitaire() -> list[str]:
    """Les FINESS certifiés au moins une fois : là où un rapport PEUT exister."""
    if not INDEX_SANITAIRE.exists():
        telecharger_index()

    classeur = openpyxl.load_workbook(INDEX_SANITAIRE, read_only=True, data_only=True)
    try:
        feuille = classeur[classeur.sheetnames[0]]
        lignes = feuille.iter_rows(values_only=True)
        position = {nom: i for i, nom in enumerate(next(lignes)) if nom}
        c_finess, c_certif = position["num_finess_et"], position["has_certif"]

        vus, sortie = set(), []
        for ligne in lignes:
            if not ligne[c_certif] or ligne[c_finess] is None:
                continue                      # 9 532 FINESS sans fiche utile
            finess = str(ligne[c_finess]).strip()
            if finess and finess not in vus:
                vus.add(finess)
                sortie.append(finess)
        return sortie                          # 8 987 FINESS
    finally:
        classeur.close()
